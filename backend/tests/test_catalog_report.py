"""Integration coverage for generated DB analysis DOCX reports."""

from __future__ import annotations

import hashlib
import io
import os
import uuid

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.db.session import get_catalog_session_factory
from app.models.catalog import (
    CatalogAnalysisRun,
    CatalogColumn,
    CatalogIndex,
    CatalogIndexColumn,
    CatalogRelation,
    CatalogRelationColumn,
    CatalogSource,
    CatalogTable,
)
from app.models.catalog_category import CatalogCategory, CatalogTableCategory


def _catalog_env() -> None:
    os.environ.setdefault("MEDICAL_DB_HOST", os.getenv("MEDICAL_DB_HOST", "localhost"))
    os.environ.setdefault("MEDICAL_DB_PORT", os.getenv("MEDICAL_DB_PORT", "5433"))
    os.environ.setdefault("CATALOG_DB_HOST", os.getenv("CATALOG_DB_HOST", "localhost"))
    os.environ.setdefault("CATALOG_DB_PORT", os.getenv("CATALOG_DB_PORT", "5434"))


@pytest.fixture()
def client():
    _catalog_env()
    from app.core.config import get_settings
    from app.db import session as session_mod
    from app.db.catalog_bootstrap import ensure_catalog_schema
    from app.main import create_app

    get_settings.cache_clear()
    session_mod._engines.clear()
    session_mod._session_factories.clear()
    ensure_catalog_schema()
    with TestClient(create_app()) as test_client:
        yield test_client


def _fp(seed: str) -> str:
    return hashlib.sha256(seed.encode("utf-8")).hexdigest()


def _create_fixture() -> tuple[int, str, str]:
    suffix = uuid.uuid4().hex[:8]
    success_fp = _fp("report-success-" + suffix)
    failed_fp = _fp("report-failed-" + suffix)
    session = get_catalog_session_factory()()
    try:
        source = CatalogSource(
            source_name=f"report_source_{suffix}",
            db_type="oracle",
            host="SECRET_REPORT_HOST",
            port=1521,
            database_name="FREEPDB1",
            default_schema="DEMIS_OWNER",
            username="SECRET_REPORT_USER",
            encrypted_password="SECRET_REPORT_CIPHERTEXT",
            connection_options={"secret_option": "DO_NOT_REPORT"},
            enabled=True,
        )
        session.add(source)
        session.flush()

        success_run = CatalogAnalysisRun(
            source_id=source.id,
            status="SUCCESS",
            target_schema="DEMIS_OWNER",
            table_count=2,
            column_count=4,
            relation_count=2,
            index_count=1,
            schema_fingerprint=success_fp,
        )
        session.add(success_run)
        session.flush()

        # Newer FAILED run must not replace the successful catalog snapshot provenance.
        failed_run = CatalogAnalysisRun(
            source_id=source.id,
            status="FAILED",
            target_schema="DEMIS_OWNER",
            table_count=0,
            column_count=0,
            relation_count=0,
            index_count=0,
            schema_fingerprint=failed_fp,
            error_message="synthetic failure",
        )
        session.add(failed_run)
        session.flush()

        patient = CatalogTable(
            source_id=source.id,
            schema_name="DEMIS_OWNER",
            table_name="TB_PT_MST",
            table_type="BASE TABLE",
            table_comment="환자 기본정보",
            object_fingerprint=_fp("report-patient-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        encounter = CatalogTable(
            source_id=source.id,
            schema_name="DEMIS_OWNER",
            table_name="TB_ENC_HIST",
            table_type="BASE TABLE",
            table_comment="진료 이력",
            object_fingerprint=_fp("report-encounter-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        session.add_all([patient, encounter])
        session.flush()

        patient_id = CatalogColumn(
            table_id=patient.id,
            ordinal_position=1,
            column_name="PT_ID",
            data_type="NUMBER",
            numeric_precision=18,
            numeric_scale=0,
            is_nullable=False,
            column_comment="환자 식별자",
            is_primary_key=True,
            is_unique=True,
            object_fingerprint=_fp("report-pt-id-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        parent_patient_id = CatalogColumn(
            table_id=patient.id,
            ordinal_position=2,
            column_name="PARENT_PT_ID",
            data_type="NUMBER",
            numeric_precision=18,
            numeric_scale=0,
            is_nullable=True,
            column_comment=None,
            is_primary_key=False,
            is_unique=False,
            object_fingerprint=_fp("report-parent-pt-id-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        encounter_id = CatalogColumn(
            table_id=encounter.id,
            ordinal_position=1,
            column_name="ENC_ID",
            data_type="NUMBER",
            is_nullable=False,
            column_comment="진료 식별자",
            is_primary_key=True,
            is_unique=True,
            object_fingerprint=_fp("report-enc-id-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        encounter_patient_id = CatalogColumn(
            table_id=encounter.id,
            ordinal_position=2,
            column_name="PT_ID",
            data_type="NUMBER",
            is_nullable=False,
            column_comment="환자 식별자",
            is_primary_key=False,
            is_unique=False,
            object_fingerprint=_fp("report-enc-pt-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        session.add_all(
            [patient_id, parent_patient_id, encounter_id, encounter_patient_id]
        )
        session.flush()

        relation = CatalogRelation(
            source_id=source.id,
            constraint_name="FK_ENC_PATIENT",
            source_table_id=encounter.id,
            target_table_id=patient.id,
            relation_type="FOREIGN_KEY",
            object_fingerprint=_fp("report-relation-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        self_relation = CatalogRelation(
            source_id=source.id,
            constraint_name="FK_PT_PARENT",
            source_table_id=patient.id,
            target_table_id=patient.id,
            relation_type="FOREIGN_KEY",
            object_fingerprint=_fp("report-self-relation-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        session.add_all([relation, self_relation])
        session.flush()
        session.add_all(
            [
                CatalogRelationColumn(
                    relation_id=relation.id,
                    ordinal_position=1,
                    source_column_id=encounter_patient_id.id,
                    target_column_id=patient_id.id,
                ),
                CatalogRelationColumn(
                    relation_id=self_relation.id,
                    ordinal_position=1,
                    source_column_id=parent_patient_id.id,
                    target_column_id=patient_id.id,
                ),
            ]
        )

        index = CatalogIndex(
            table_id=encounter.id,
            index_name="IX_ENC_PT_ID",
            is_unique=False,
            index_method="BTREE",
            index_definition=None,
            object_fingerprint=_fp("report-index-" + suffix),
            last_run_id=success_run.id,
            active=True,
        )
        session.add(index)
        session.flush()
        session.add(
            CatalogIndexColumn(
                index_id=index.id,
                ordinal_position=1,
                column_name="PT_ID",
            )
        )

        category = CatalogCategory(
            source_id=source.id,
            category_key="patient",
            category_name="환자",
            description="환자 관련 테이블",
            sort_order=10,
            active=True,
        )
        session.add(category)
        session.flush()
        session.add(
            CatalogTableCategory(
                table_id=patient.id,
                category_id=category.id,
                is_primary=True,
                assignment_source="MANUAL",
                confidence=None,
                note="검토 완료",
            )
        )

        session.commit()
        return int(source.id), success_fp, failed_fp
    finally:
        session.close()


def _create_multi_schema_fixture() -> int:
    suffix = uuid.uuid4().hex[:8]
    session = get_catalog_session_factory()()
    try:
        source = CatalogSource(
            source_name=f"report_multi_{suffix}",
            db_type="oracle",
            host="localhost",
            port=1521,
            database_name="FREEPDB1",
            default_schema="DEMIS_OWNER",
            username="fixture",
            encrypted_password=None,
            connection_options=None,
            enabled=True,
        )
        session.add(source)
        session.flush()
        run = CatalogAnalysisRun(
            source_id=source.id,
            status="SUCCESS",
            target_schema="DEMIS_OWNER",
            table_count=2,
            column_count=2,
            relation_count=1,
            index_count=0,
            schema_fingerprint=_fp("report-multi-run-" + suffix),
        )
        session.add(run)
        session.flush()

        left = CatalogTable(
            source_id=source.id,
            schema_name="DEMIS_OWNER",
            table_name="TB_LEFT",
            table_type="BASE TABLE",
            table_comment=None,
            object_fingerprint=_fp("report-left-" + suffix),
            last_run_id=run.id,
            active=True,
        )
        right = CatalogTable(
            source_id=source.id,
            schema_name="REF_OWNER",
            table_name="TB_RIGHT",
            table_type="BASE TABLE",
            table_comment=None,
            object_fingerprint=_fp("report-right-" + suffix),
            last_run_id=run.id,
            active=True,
        )
        session.add_all([left, right])
        session.flush()

        left_id = CatalogColumn(
            table_id=left.id,
            ordinal_position=1,
            column_name="RIGHT_ID",
            data_type="NUMBER",
            is_nullable=False,
            column_comment=None,
            is_primary_key=False,
            is_unique=False,
            object_fingerprint=_fp("report-left-id-" + suffix),
            last_run_id=run.id,
            active=True,
        )
        right_id = CatalogColumn(
            table_id=right.id,
            ordinal_position=1,
            column_name="RIGHT_ID",
            data_type="NUMBER",
            is_nullable=False,
            column_comment=None,
            is_primary_key=True,
            is_unique=True,
            object_fingerprint=_fp("report-right-id-" + suffix),
            last_run_id=run.id,
            active=True,
        )
        session.add_all([left_id, right_id])
        session.flush()

        relation = CatalogRelation(
            source_id=source.id,
            constraint_name="FK_LEFT_RIGHT",
            source_table_id=left.id,
            target_table_id=right.id,
            relation_type="FOREIGN_KEY",
            object_fingerprint=_fp("report-multi-relation-" + suffix),
            last_run_id=run.id,
            active=True,
        )
        session.add(relation)
        session.flush()
        session.add(
            CatalogRelationColumn(
                relation_id=relation.id,
                ordinal_position=1,
                source_column_id=left_id.id,
                target_column_id=right_id.id,
            )
        )
        session.commit()
        return int(source.id)
    finally:
        session.close()


def _cleanup(source_id: int) -> None:
    session = get_catalog_session_factory()()
    try:
        session.execute(
            text(
                "DELETE FROM catalog_table_category WHERE table_id IN "
                "(SELECT id FROM catalog_table WHERE source_id = :sid)"
            ),
            {"sid": source_id},
        )
        session.execute(
            text("DELETE FROM catalog_category WHERE source_id = :sid"),
            {"sid": source_id},
        )
        session.execute(
            text("DELETE FROM catalog_relation WHERE source_id = :sid"),
            {"sid": source_id},
        )
        session.execute(
            text(
                "DELETE FROM catalog_index WHERE table_id IN "
                "(SELECT id FROM catalog_table WHERE source_id = :sid)"
            ),
            {"sid": source_id},
        )
        session.execute(
            text(
                "DELETE FROM catalog_column WHERE table_id IN "
                "(SELECT id FROM catalog_table WHERE source_id = :sid)"
            ),
            {"sid": source_id},
        )
        session.execute(
            text("UPDATE catalog_table SET last_run_id = NULL WHERE source_id = :sid"),
            {"sid": source_id},
        )
        session.execute(
            text("DELETE FROM catalog_table WHERE source_id = :sid"),
            {"sid": source_id},
        )
        session.execute(
            text("DELETE FROM catalog_analysis_run WHERE source_id = :sid"),
            {"sid": source_id},
        )
        session.execute(
            text("DELETE FROM catalog_source WHERE id = :sid"),
            {"sid": source_id},
        )
        session.commit()
    except Exception:  # noqa: BLE001
        session.rollback()
        raise
    finally:
        session.close()


def _document_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _document_tables(content: bytes) -> list[list[list[str]]]:
    document = Document(io.BytesIO(content))
    return [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]


def test_db_analysis_report_download_uses_latest_success_and_excludes_secrets(
    client: TestClient,
) -> None:
    source_id, success_fp, failed_fp = _create_fixture()
    try:
        metadata_response = client.get(f"/api/v1/catalog/report/{source_id}/metadata")
        assert metadata_response.status_code == 200, metadata_response.text
        metadata = metadata_response.json()
        assert metadata["counts"]["tables"] == 2
        assert metadata["counts"]["relations"] == 2
        assert metadata["counts"]["categories"] == 1
        assert metadata["schema_fingerprint"] == success_fp
        assert metadata["schema_fingerprint"] != failed_fp
        assert metadata["security"]["credentials_exported"] is False

        response = client.get(f"/api/v1/catalog/report/{source_id}/download")
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "attachment; filename=" in response.headers["content-disposition"]
        assert response.headers["x-catalog-report-version"] == "1.0"

        text_content = _document_text(response.content)
        assert "DEMIS DB 분석서" in text_content
        assert "TB_PT_MST" in text_content
        assert "TB_ENC_HIST" in text_content
        assert "FK_ENC_PATIENT" in text_content
        assert "FK_PT_PARENT" in text_content
        assert "IX_ENC_PT_ID" in text_content
        assert "환자 기본정보" in text_content
        assert "MANUAL" in text_content
        assert "자기참조 FK는 상세 관계에서 SELF로 표시" in text_content
        assert success_fp in text_content
        assert failed_fp not in text_content

        tables = _document_tables(response.content)

        table_summary = next(
            table
            for table in tables
            if table and table[0] == ["#", "Table", "Type", "Columns", "PK", "Comment", "Category"]
        )
        assert all("Schema" not in cell for cell in table_summary[0])
        assert {row[1] for row in table_summary[1:]} == {"TB_ENC_HIST", "TB_PT_MST"}

        relation_summary = next(
            table
            for table in tables
            if table
            and table[0] == ["Constraint", "Source", "Target", "Type", "Column Mapping"]
        )
        enc_relation = next(row for row in relation_summary[1:] if row[0] == "FK_ENC_PATIENT")
        assert enc_relation[1] == "TB_ENC_HIST"
        assert enc_relation[2] == "TB_PT_MST"
        self_summary = next(row for row in relation_summary[1:] if row[0] == "FK_PT_PARENT")
        assert self_summary[1] == "TB_PT_MST"
        assert self_summary[2] == "TB_PT_MST"

        index_summary = next(
            table
            for table in tables
            if table and table[0] == ["Table", "Index", "Unique", "Method", "Columns"]
        )
        index_row = next(row for row in index_summary[1:] if row[1] == "IX_ENC_PT_ID")
        assert index_row[0] == "TB_ENC_HIST"

        category_assignments = next(
            table
            for table in tables
            if table
            and table[0] == ["Table", "Category", "Primary", "Source", "Confidence", "Note"]
        )
        assert category_assignments[1][0] == "TB_PT_MST"

        self_detail_rows = [
            row
            for table in tables
            if table and table[0] == ["Direction", "Constraint", "Related Table", "Mapping"]
            for row in table[1:]
            if row[1] == "FK_PT_PARENT"
        ]
        assert self_detail_rows == [
            ["SELF", "FK_PT_PARENT", "DEMIS_OWNER.TB_PT_MST", "PARENT_PT_ID → PT_ID"]
        ]

        for secret in [
            "SECRET_REPORT_HOST",
            "SECRET_REPORT_USER",
            "SECRET_REPORT_CIPHERTEXT",
            "DO_NOT_REPORT",
        ]:
            assert secret not in text_content
    finally:
        _cleanup(source_id)


def test_db_analysis_report_multi_schema_keeps_qualified_summary_names(
    client: TestClient,
) -> None:
    source_id = _create_multi_schema_fixture()
    try:
        response = client.get(f"/api/v1/catalog/report/{source_id}/download")
        assert response.status_code == 200, response.text
        tables = _document_tables(response.content)

        table_summary = next(
            table
            for table in tables
            if table
            and table[0]
            == ["#", "Schema", "Table", "Type", "Columns", "PK", "Comment", "Category"]
        )
        assert {row[1] for row in table_summary[1:]} == {"DEMIS_OWNER", "REF_OWNER"}

        relation_summary = next(
            table
            for table in tables
            if table
            and table[0] == ["Constraint", "Source", "Target", "Type", "Column Mapping"]
        )
        relation_row = next(row for row in relation_summary[1:] if row[0] == "FK_LEFT_RIGHT")
        assert relation_row[1] == "DEMIS_OWNER.TB_LEFT"
        assert relation_row[2] == "REF_OWNER.TB_RIGHT"
    finally:
        _cleanup(source_id)


def test_db_analysis_report_unknown_source_returns_404(client: TestClient) -> None:
    response = client.get("/api/v1/catalog/report/999999999/download")
    assert response.status_code == 404
    assert response.json()["detail"] == "catalog source not found"
