"""Generate a human-readable DB analysis report from catalog metadata."""

from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.catalog import (
    CatalogAnalysisRun,
    CatalogColumn,
    CatalogIndex,
    CatalogIndexColumn,
    CatalogKeyConstraint,
    CatalogKeyConstraintColumn,
    CatalogRelation,
    CatalogRelationColumn,
    CatalogSource,
    CatalogTable,
)
from app.models.catalog_category import CatalogCategory, CatalogTableCategory


REPORT_VERSION = "1.0"


@dataclass(frozen=True)
class CatalogReportResult:
    filename: str
    content: bytes
    metadata: dict[str, Any]


def _iso(value: datetime | None) -> str | None:
    return value.isoformat() if value is not None else None


def _safe_text(value: Any, empty: str = "-") -> str:
    if value is None:
        return empty
    text = str(value)
    return text if text else empty


def _table_key(table: CatalogTable) -> str:
    return f"{table.schema_name}.{table.table_name}"


def _single_schema_name(tables: list[CatalogTable]) -> str | None:
    schema_names = {str(table.schema_name) for table in tables if table.schema_name}
    return next(iter(schema_names)) if len(schema_names) == 1 else None


def _summary_table_name(table: CatalogTable, single_schema: str | None) -> str:
    return table.table_name if single_schema else _table_key(table)


def _format_data_type(column: CatalogColumn) -> str:
    base = _safe_text(column.data_type)
    if column.character_maximum_length is not None:
        return f"{base}({column.character_maximum_length})"
    if column.numeric_precision is not None:
        if column.numeric_scale is not None:
            return f"{base}({column.numeric_precision},{column.numeric_scale})"
        return f"{base}({column.numeric_precision})"
    return base


def _set_cell_text(
    cell,
    text: Any,
    *,
    bold: bool = False,
    font_size: float = 8.5,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(_safe_text(text))
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
    *,
    font_size: float = 8.0,
):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    _set_repeat_table_header(header)
    for idx, header_text in enumerate(headers):
        _set_cell_text(header.cells[idx], header_text, bold=True, font_size=font_size)
        _shade_cell(header.cells[idx], "D9EAF7")
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            _set_cell_text(row.cells[idx], value, font_size=font_size)
    return table


def _set_doc_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    for style_name, size in [
        ("Title", 22),
        ("Heading 1", 16),
        ("Heading 2", 13),
        ("Heading 3", 11),
    ]:
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style.font.size = Pt(size)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def _add_cover(
    document: Document,
    source: CatalogSource,
    latest_success: CatalogAnalysisRun | None,
) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_after = Pt(18)
    run = p.add_run("DEMIS DB 분석서")
    run.bold = True
    run.font.size = Pt(24)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(_safe_text(source.source_name))
    run.font.size = Pt(14)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    document.add_paragraph("")
    overview_rows = [
        ["DBMS", source.db_type],
        ["Database", source.database_name],
        ["Default Schema", source.default_schema],
        ["Catalog Source", source.source_name],
        ["최근 성공 분석 Run", latest_success.id if latest_success else "-"],
        ["Schema Fingerprint", latest_success.schema_fingerprint if latest_success else "-"],
        ["보고서 버전", REPORT_VERSION],
        ["생성시각(UTC)", datetime.now(timezone.utc).isoformat()],
    ]
    _add_table(document, ["항목", "값"], overview_rows, font_size=9)

    document.add_paragraph("")
    note = document.add_paragraph()
    run = note.add_run(
        "본 문서는 DEMIS Schema Analyzer가 수집한 Catalog 메타데이터를 기반으로 자동 생성되었습니다. "
        "DB COMMENT 등 물리 DB에서 확인된 정보와 Category 등 의미 메타데이터를 구분하여 표기하며, "
        "확인되지 않은 업무 의미를 임의로 생성하지 않습니다."
    )
    run.font.size = Pt(9)
    document.add_page_break()


def build_db_analysis_report(session: Session, source_id: int) -> CatalogReportResult:
    source = session.get(CatalogSource, source_id)
    if source is None:
        raise LookupError("catalog source not found")

    latest_success = session.scalar(
        select(CatalogAnalysisRun)
        .where(
            CatalogAnalysisRun.source_id == source_id,
            CatalogAnalysisRun.status == "SUCCESS",
        )
        .order_by(CatalogAnalysisRun.id.desc())
        .limit(1)
    )

    tables = session.scalars(
        select(CatalogTable)
        .where(CatalogTable.source_id == source_id, CatalogTable.active.is_(True))
        .order_by(CatalogTable.schema_name, CatalogTable.table_name)
    ).all()
    table_by_id = {int(table.id): table for table in tables}
    table_ids = list(table_by_id)
    single_schema = _single_schema_name(tables)

    columns: list[CatalogColumn] = []
    indexes: list[CatalogIndex] = []
    relations: list[CatalogRelation] = []
    constraints: list[CatalogKeyConstraint] = []
    if table_ids:
        columns = session.scalars(
            select(CatalogColumn)
            .where(CatalogColumn.table_id.in_(table_ids), CatalogColumn.active.is_(True))
            .order_by(CatalogColumn.table_id, CatalogColumn.ordinal_position)
        ).all()
        indexes = session.scalars(
            select(CatalogIndex)
            .where(CatalogIndex.table_id.in_(table_ids), CatalogIndex.active.is_(True))
            .order_by(CatalogIndex.table_id, CatalogIndex.index_name)
        ).all()
        relations = session.scalars(
            select(CatalogRelation)
            .where(CatalogRelation.source_id == source_id, CatalogRelation.active.is_(True))
            .order_by(CatalogRelation.constraint_name, CatalogRelation.id)
        ).all()
        constraints = session.scalars(
            select(CatalogKeyConstraint)
            .where(
                CatalogKeyConstraint.table_id.in_(table_ids),
                CatalogKeyConstraint.active.is_(True),
            )
            .order_by(CatalogKeyConstraint.table_id, CatalogKeyConstraint.constraint_name)
        ).all()

    columns_by_table: dict[int, list[CatalogColumn]] = defaultdict(list)
    for column in columns:
        columns_by_table[int(column.table_id)].append(column)
    column_by_id = {int(column.id): column for column in columns}

    index_columns: dict[int, list[CatalogIndexColumn]] = defaultdict(list)
    index_ids = [int(index.id) for index in indexes]
    if index_ids:
        rows = session.scalars(
            select(CatalogIndexColumn)
            .where(CatalogIndexColumn.index_id.in_(index_ids))
            .order_by(CatalogIndexColumn.index_id, CatalogIndexColumn.ordinal_position)
        ).all()
        for row in rows:
            index_columns[int(row.index_id)].append(row)

    indexes_by_table: dict[int, list[CatalogIndex]] = defaultdict(list)
    for index in indexes:
        indexes_by_table[int(index.table_id)].append(index)

    constraint_columns: dict[int, list[CatalogKeyConstraintColumn]] = defaultdict(list)
    constraint_ids = [int(item.id) for item in constraints]
    if constraint_ids:
        rows = session.scalars(
            select(CatalogKeyConstraintColumn)
            .where(CatalogKeyConstraintColumn.constraint_id.in_(constraint_ids))
            .order_by(
                CatalogKeyConstraintColumn.constraint_id,
                CatalogKeyConstraintColumn.ordinal_position,
            )
        ).all()
        for row in rows:
            constraint_columns[int(row.constraint_id)].append(row)

    constraints_by_table: dict[int, list[CatalogKeyConstraint]] = defaultdict(list)
    for constraint in constraints:
        constraints_by_table[int(constraint.table_id)].append(constraint)

    relation_columns: dict[int, list[CatalogRelationColumn]] = defaultdict(list)
    relation_ids = [int(item.id) for item in relations]
    if relation_ids:
        rows = session.scalars(
            select(CatalogRelationColumn)
            .where(CatalogRelationColumn.relation_id.in_(relation_ids))
            .order_by(
                CatalogRelationColumn.relation_id,
                CatalogRelationColumn.ordinal_position,
            )
        ).all()
        for row in rows:
            relation_columns[int(row.relation_id)].append(row)

    inbound_by_table: dict[int, list[CatalogRelation]] = defaultdict(list)
    outbound_by_table: dict[int, list[CatalogRelation]] = defaultdict(list)
    relation_rows: list[list[Any]] = []
    for relation in relations:
        source_table = table_by_id.get(int(relation.source_table_id))
        target_table = table_by_id.get(int(relation.target_table_id))
        if source_table is None or target_table is None:
            continue
        outbound_by_table[int(source_table.id)].append(relation)
        inbound_by_table[int(target_table.id)].append(relation)
        mapping = []
        for item in relation_columns.get(int(relation.id), []):
            source_column = column_by_id.get(int(item.source_column_id))
            target_column = column_by_id.get(int(item.target_column_id))
            if source_column is not None and target_column is not None:
                mapping.append(f"{source_column.column_name} → {target_column.column_name}")
        relation_rows.append(
            [
                relation.constraint_name,
                _summary_table_name(source_table, single_schema),
                _summary_table_name(target_table, single_schema),
                relation.relation_type,
                ", ".join(mapping) or "-",
            ]
        )

    categories = session.scalars(
        select(CatalogCategory)
        .where(CatalogCategory.source_id == source_id)
        .order_by(CatalogCategory.sort_order, CatalogCategory.category_name)
    ).all()
    category_by_id = {int(category.id): category for category in categories}
    mappings: list[CatalogTableCategory] = []
    if table_ids and category_by_id:
        mappings = session.scalars(
            select(CatalogTableCategory)
            .where(
                CatalogTableCategory.table_id.in_(table_ids),
                CatalogTableCategory.category_id.in_(list(category_by_id)),
            )
            .order_by(
                CatalogTableCategory.table_id,
                CatalogTableCategory.is_primary.desc(),
            )
        ).all()

    categories_by_table: dict[int, list[CatalogTableCategory]] = defaultdict(list)
    for mapping in mappings:
        categories_by_table[int(mapping.table_id)].append(mapping)

    document = Document()
    _set_doc_defaults(document)
    _add_cover(document, source, latest_success)

    document.add_heading("1. 분석 개요", level=1)
    summary_rows = [
        ["Catalog Source", source.source_name],
        ["DBMS", source.db_type],
        ["Database", source.database_name],
        ["Default Schema", source.default_schema],
        ["최근 성공 분석 Run", latest_success.id if latest_success else "-"],
        ["최근 성공 분석 상태", latest_success.status if latest_success else "-"],
        [
            "분석 대상 Schema",
            latest_success.target_schema if latest_success else source.default_schema,
        ],
        ["Schema Fingerprint", latest_success.schema_fingerprint if latest_success else "-"],
        ["분석 시작", _iso(latest_success.started_at) if latest_success else "-"],
        ["분석 종료", _iso(latest_success.finished_at) if latest_success else "-"],
    ]
    _add_table(document, ["항목", "값"], summary_rows, font_size=8.5)

    document.add_heading("1.1 Catalog 현황", level=2)
    table_comment_count = sum(1 for table in tables if table.table_comment)
    column_comment_count = sum(1 for column in columns if column.column_comment)
    count_rows = [
        ["Tables", len(tables)],
        ["Columns", len(columns)],
        ["Relations(FK)", len(relation_rows)],
        ["Indexes", len(indexes)],
        ["Key Constraints", len(constraints)],
        ["Categories", len(categories)],
        ["Category Assignments", len(mappings)],
        ["Table DB COMMENT", f"{table_comment_count}/{len(tables)}"],
        ["Column DB COMMENT", f"{column_comment_count}/{len(columns)}"],
    ]
    _add_table(document, ["구분", "수량/현황"], count_rows, font_size=8.5)

    document.add_heading("1.2 정보 출처 및 판정 기준", level=2)
    for text in [
        (
            "Table·Column·PK/FK·Index 등 물리 구조는 Schema Analyzer가 수집한 "
            "Catalog 값을 사용합니다."
        ),
        (
            "Table/Column 설명은 DB COMMENT가 실제 존재하는 경우에만 "
            "DB_COMMENT로 표기합니다."
        ),
        (
            "Category는 물리 DB 사실과 분리된 의미 메타데이터이며, "
            "MANUAL/AUTO/IMPORT assignment_source를 유지합니다."
        ),
        (
            "DB COMMENT가 없는 항목에 대해 Table명/Column명만으로 업무 의미를 "
            "자동 생성하지 않습니다."
        ),
        (
            "자기참조 FK는 상세 관계에서 SELF로 표시하며 Outbound/Inbound "
            "집계에는 각각 포함합니다."
        ),
        (
            "단일 Schema 보고서의 요약표에서는 가독성을 위해 Schema 접두어를 "
            "생략하고 상세 명세에서는 전체 Schema.Table을 유지합니다."
        ),
        (
            "보고서에는 접속 Host, Username, Password, Connection Option 등 "
            "Target 연결 비밀정보를 포함하지 않습니다."
        ),
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("2. Table 요약", level=1)
    table_summary_rows: list[list[Any]] = []
    for no, table in enumerate(tables, start=1):
        pk_columns = [
            column.column_name
            for column in columns_by_table.get(int(table.id), [])
            if bool(column.is_primary_key)
        ]
        category_names = []
        for mapping in categories_by_table.get(int(table.id), []):
            category = category_by_id.get(int(mapping.category_id))
            if category is not None:
                suffix = "*" if mapping.is_primary else ""
                category_names.append(f"{category.category_name}{suffix}")
        comment = (
            f"[DB_COMMENT] {table.table_comment}" if table.table_comment else "-"
        )
        if single_schema:
            table_summary_rows.append(
                [
                    no,
                    table.table_name,
                    table.table_type,
                    len(columns_by_table.get(int(table.id), [])),
                    ", ".join(pk_columns) or "-",
                    comment,
                    ", ".join(category_names) or "-",
                ]
            )
        else:
            table_summary_rows.append(
                [
                    no,
                    table.schema_name,
                    table.table_name,
                    table.table_type,
                    len(columns_by_table.get(int(table.id), [])),
                    ", ".join(pk_columns) or "-",
                    comment,
                    ", ".join(category_names) or "-",
                ]
            )
    if single_schema:
        document.add_paragraph(f"Schema: {single_schema} (단일 Schema)")
        table_headers = [
            "#",
            "Table",
            "Type",
            "Columns",
            "PK",
            "Comment",
            "Category",
        ]
    else:
        table_headers = [
            "#",
            "Schema",
            "Table",
            "Type",
            "Columns",
            "PK",
            "Comment",
            "Category",
        ]
    _add_table(document, table_headers, table_summary_rows, font_size=7.2)

    document.add_heading("3. 관계(FK) 요약", level=1)
    if single_schema:
        document.add_paragraph(
            f"Schema: {single_schema} (단일 Schema, Source/Target 접두어 생략)"
        )
    if relation_rows:
        _add_table(
            document,
            ["Constraint", "Source", "Target", "Type", "Column Mapping"],
            relation_rows,
            font_size=7.3,
        )
    else:
        document.add_paragraph("활성 FK 관계가 없습니다.")

    document.add_heading("4. Index 요약", level=1)
    if single_schema:
        document.add_paragraph(
            f"Schema: {single_schema} (단일 Schema, Table 접두어 생략)"
        )
    index_rows: list[list[Any]] = []
    for index in indexes:
        table = table_by_id.get(int(index.table_id))
        if table is None:
            continue
        index_rows.append(
            [
                _summary_table_name(table, single_schema),
                index.index_name,
                "Y" if index.is_unique else "N",
                index.index_method or "-",
                ", ".join(
                    item.column_name
                    for item in index_columns.get(int(index.id), [])
                )
                or "-",
            ]
        )
    if index_rows:
        _add_table(
            document,
            ["Table", "Index", "Unique", "Method", "Columns"],
            index_rows,
            font_size=7.6,
        )
    else:
        document.add_paragraph("활성 Index가 없습니다.")

    document.add_heading("5. Category / 의미 메타데이터", level=1)
    if categories:
        category_rows = [
            [
                category.category_key,
                category.category_name,
                category.description or "-",
                "Y" if category.active else "N",
            ]
            for category in categories
        ]
        _add_table(
            document,
            ["Key", "Category", "Description", "Active"],
            category_rows,
            font_size=8,
        )

        document.add_heading("5.1 Table Category 할당", level=2)
        assignment_rows: list[list[Any]] = []
        for mapping in mappings:
            table = table_by_id.get(int(mapping.table_id))
            category = category_by_id.get(int(mapping.category_id))
            if table is None or category is None:
                continue
            assignment_rows.append(
                [
                    _summary_table_name(table, single_schema),
                    category.category_name,
                    "Y" if mapping.is_primary else "N",
                    mapping.assignment_source,
                    mapping.confidence if mapping.confidence is not None else "-",
                    mapping.note or "-",
                ]
            )
        if single_schema and assignment_rows:
            document.add_paragraph(
                f"Schema: {single_schema} (단일 Schema, Table 접두어 생략)"
            )
        if assignment_rows:
            _add_table(
                document,
                ["Table", "Category", "Primary", "Source", "Confidence", "Note"],
                assignment_rows,
                font_size=7.5,
            )
        else:
            document.add_paragraph("Table Category 할당이 없습니다.")
    else:
        document.add_paragraph("등록된 Category가 없습니다.")

    document.add_page_break()
    document.add_heading("6. Table 상세 명세", level=1)
    for table_idx, table in enumerate(tables, start=1):
        if table_idx > 1:
            document.add_page_break()
        document.add_heading(f"6.{table_idx} {_table_key(table)}", level=2)

        category_texts = []
        for mapping in categories_by_table.get(int(table.id), []):
            category = category_by_id.get(int(mapping.category_id))
            if category is None:
                continue
            category_texts.append(
                f"{category.category_name}"
                f"(source={mapping.assignment_source}, "
                f"confidence={mapping.confidence if mapping.confidence is not None else '-'})"
            )

        table_meta_rows = [
            ["Schema", table.schema_name],
            ["Table", table.table_name],
            ["Type", table.table_type],
            [
                "Comment",
                f"[DB_COMMENT] {table.table_comment}"
                if table.table_comment
                else "-",
            ],
            ["Object Fingerprint", table.object_fingerprint],
            ["Categories", "; ".join(category_texts) or "-"],
            ["Outbound FK", len(outbound_by_table.get(int(table.id), []))],
            ["Inbound FK", len(inbound_by_table.get(int(table.id), []))],
            ["Indexes", len(indexes_by_table.get(int(table.id), []))],
        ]
        _add_table(document, ["항목", "값"], table_meta_rows, font_size=8.2)

        document.add_heading("Columns", level=3)
        detail_column_rows: list[list[Any]] = []
        for column in columns_by_table.get(int(table.id), []):
            flags = []
            if column.is_primary_key:
                flags.append("PK")
            if column.is_unique:
                flags.append("UNIQUE")
            if not column.is_nullable:
                flags.append("NOT NULL")
            detail_column_rows.append(
                [
                    column.ordinal_position,
                    column.column_name,
                    _format_data_type(column),
                    ", ".join(flags) or "-",
                    column.default_value or "-",
                    f"[DB_COMMENT] {column.column_comment}"
                    if column.column_comment
                    else "-",
                ]
            )
        if detail_column_rows:
            _add_table(
                document,
                ["#", "Column", "Type", "Flags", "Default", "Comment"],
                detail_column_rows,
                font_size=7.3,
            )
        else:
            document.add_paragraph("활성 Column이 없습니다.")

        table_constraints = constraints_by_table.get(int(table.id), [])
        if table_constraints:
            document.add_heading("Key Constraints", level=3)
            constraint_rows = []
            for constraint in table_constraints:
                names = []
                for item in constraint_columns.get(int(constraint.id), []):
                    column = column_by_id.get(int(item.column_id))
                    if column is not None:
                        names.append(column.column_name)
                constraint_rows.append(
                    [
                        constraint.constraint_name,
                        constraint.constraint_type,
                        ", ".join(names) or "-",
                    ]
                )
            _add_table(
                document,
                ["Constraint", "Type", "Columns"],
                constraint_rows,
                font_size=7.8,
            )

        table_relations = (
            outbound_by_table.get(int(table.id), [])
            + inbound_by_table.get(int(table.id), [])
        )
        if table_relations:
            document.add_heading("Relations", level=3)
            detail_relation_rows = []
            seen: set[int] = set()
            for relation in table_relations:
                if int(relation.id) in seen:
                    continue
                seen.add(int(relation.id))
                source_table = table_by_id.get(int(relation.source_table_id))
                target_table = table_by_id.get(int(relation.target_table_id))
                if source_table is None or target_table is None:
                    continue
                mapping = []
                for item in relation_columns.get(int(relation.id), []):
                    source_column = column_by_id.get(int(item.source_column_id))
                    target_column = column_by_id.get(int(item.target_column_id))
                    if source_column is not None and target_column is not None:
                        mapping.append(
                            f"{source_column.column_name} → {target_column.column_name}"
                        )

                is_self = (
                    int(relation.source_table_id) == int(table.id)
                    and int(relation.target_table_id) == int(table.id)
                )
                if is_self:
                    direction = "SELF"
                    peer = table
                elif int(relation.source_table_id) == int(table.id):
                    direction = "OUT"
                    peer = target_table
                else:
                    direction = "IN"
                    peer = source_table

                detail_relation_rows.append(
                    [
                        direction,
                        relation.constraint_name,
                        _table_key(peer),
                        ", ".join(mapping) or "-",
                    ]
                )
            _add_table(
                document,
                ["Direction", "Constraint", "Related Table", "Mapping"],
                detail_relation_rows,
                font_size=7.5,
            )

        table_indexes = indexes_by_table.get(int(table.id), [])
        if table_indexes:
            document.add_heading("Indexes", level=3)
            detail_index_rows = []
            for index in table_indexes:
                detail_index_rows.append(
                    [
                        index.index_name,
                        "Y" if index.is_unique else "N",
                        index.index_method or "-",
                        ", ".join(
                            item.column_name
                            for item in index_columns.get(int(index.id), [])
                        )
                        or "-",
                    ]
                )
            _add_table(
                document,
                ["Index", "Unique", "Method", "Columns"],
                detail_index_rows,
                font_size=7.8,
            )

    buffer = io.BytesIO()
    document.save(buffer)

    safe_source = "".join(
        ch if ch.isalnum() or ch in {"-", "_"} else "_"
        for ch in source.source_name
    )
    filename = f"DEMIS_DB_분석서_{safe_source}.docx"
    metadata = {
        "report_version": REPORT_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": {
            "source_name": source.source_name,
            "db_type": source.db_type,
            "database_name": source.database_name,
            "default_schema": source.default_schema,
        },
        "latest_success_run_id": int(latest_success.id) if latest_success else None,
        "schema_fingerprint": (
            latest_success.schema_fingerprint if latest_success else None
        ),
        "counts": {
            "tables": len(tables),
            "columns": len(columns),
            "relations": len(relation_rows),
            "indexes": len(indexes),
            "categories": len(categories),
            "category_assignments": len(mappings),
        },
        "security": {
            "connection_host_exported": False,
            "username_exported": False,
            "credentials_exported": False,
            "connection_options_exported": False,
        },
        "provenance_note": (
            "Physical DB facts are reported from Catalog metadata. "
            "DB comments are marked as DB_COMMENT. "
            "Category assignments preserve MANUAL/AUTO/IMPORT provenance."
        ),
    }
    return CatalogReportResult(
        filename=filename,
        content=buffer.getvalue(),
        metadata=metadata,
    )
